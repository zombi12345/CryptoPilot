import logging
import uuid
from pathlib import Path
from pypdf import PdfReader
from docx import Document
import chromadb
from sentence_transformers import SentenceTransformer

logger = logging.getLogger(__name__)

CHROMA_DIR = Path("data/chroma")
CHROMA_DIR.mkdir(parents=True, exist_ok=True)

_client = chromadb.PersistentClient(path=str(CHROMA_DIR))
_model = None


def get_model():
    global _model
    if _model is None:
        _model = SentenceTransformer("sergeyzh/rubert-mini-retriever")
    return _model


def extract_text_pdf(file_bytes: bytes) -> str:
    import io
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts)


def extract_text_docx(file_bytes: bytes) -> str:
    import io
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    text = " ".join(text.split())
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap
        if start < 0:
            start = 0
        if end >= len(text):
            break
    return [c for c in chunks if len(c) > 50]


def split_into_chapters(text: str, max_chapter_chars: int = 50000) -> list:
    import re
    lines = text.split("\n")
    chapter_starts = []
    pattern = re.compile(r"^\s*(Глава|Chapter|CHAPTER|ГЛАВА)\s+[\dIVXLCDM]+", re.IGNORECASE)

    for i, line in enumerate(lines):
        if pattern.match(line) or (line.strip().isupper() and 3 < len(line.strip()) < 100):
            chapter_starts.append(i)

    if len(chapter_starts) < 2:
        chapters = []
        buf = ""
        for line in lines:
            buf += line + "\n"
            if len(buf) >= max_chapter_chars:
                chapters.append(buf)
                buf = ""
        if buf.strip():
            chapters.append(buf)
        return chapters

    chapters = []
    for idx, start in enumerate(chapter_starts):
        end = chapter_starts[idx + 1] if idx + 1 < len(chapter_starts) else len(lines)
        chunk = "\n".join(lines[start:end])
        if len(chunk) > max_chapter_chars:
            for i in range(0, len(chunk), max_chapter_chars):
                chapters.append(chunk[i:i+max_chapter_chars])
        else:
            chapters.append(chunk)
    return [c for c in chapters if len(c) > 500]


def add_book(user_id: int, filename: str, file_bytes: bytes) -> dict:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        text = extract_text_pdf(file_bytes)
    elif ext == "docx":
        text = extract_text_docx(file_bytes)
    else:
        return {"error": "Поддерживаются только PDF и DOCX"}

    if not text.strip():
        return {"error": "Не удалось извлечь текст"}

    chunks = chunk_text(text)
    if not chunks:
        return {"error": "Текст слишком короткий"}

    collection_name = f"user_{user_id}"
    collection = _client.get_or_create_collection(name=collection_name)

    model = get_model()
    embeddings = model.encode(chunks, show_progress_bar=False).tolist()

    ids = [f"{uuid.uuid4().hex}" for _ in chunks]
    metadatas = [{"source": filename, "chunk_idx": i} for i in range(len(chunks))]

    collection.add(
        ids=ids,
        embeddings=embeddings,
        documents=chunks,
        metadatas=metadatas,
    )

    return {
        "ok": True,
        "filename": filename,
        "chunks": len(chunks),
        "chars": len(text),
        "full_text": text,
    }


def list_books(user_id: int) -> list:
    collection_name = f"user_{user_id}"
    try:
        collection = _client.get_collection(name=collection_name)
    except Exception:
        return []
    data = collection.get(include=["metadatas"])
    sources = {}
    for meta in data.get("metadatas", []):
        if not meta:
            continue
        src = meta.get("source", "?")
        sources[src] = sources.get(src, 0) + 1
    return [{"filename": k, "chunks": v} for k, v in sources.items()]


def search(user_id: int, query: str, top_k: int = 20) -> list:
    collection_name = f"user_{user_id}"
    try:
        collection = _client.get_collection(name=collection_name)
    except Exception:
        return []

    model = get_model()
    query_emb = model.encode([query], show_progress_bar=False).tolist()

    results = collection.query(
        query_embeddings=query_emb,
        n_results=top_k,
        include=["documents", "metadatas", "distances"],
    )

    out = []
    docs = results.get("documents", [[]])[0]
    metas = results.get("metadatas", [[]])[0]
    dists = results.get("distances", [[]])[0]
    for doc, meta, dist in zip(docs, metas, dists):
        out.append({
            "text": doc,
            "source": meta.get("source", "?") if meta else "?",
            "score": 1 - dist if dist is not None else 0,
        })
    return out


def get_book_context(user_id: int, query: str, top_k: int = 20) -> str:
    chunks = search(user_id, query, top_k=top_k)
    if not chunks:
        return ""
    grouped = {}
    for c in chunks:
        grouped.setdefault(c["source"], []).append(c["text"])
    parts = []
    for src, texts in grouped.items():
        parts.append(f"[{src}]\n" + "\n---\n".join(texts))
    return "\n\n=====\n\n".join(parts)


def delete_book(user_id: int, filename: str) -> bool:
    collection_name = f"user_{user_id}"
    try:
        collection = _client.get_collection(name=collection_name)
    except Exception:
        return False
    data = collection.get(include=["metadatas"])
    ids_to_delete = []
    for i, meta in enumerate(data.get("metadatas", [])):
        if meta and meta.get("source") == filename:
            ids_to_delete.append(data["ids"][i])
    if ids_to_delete:
        collection.delete(ids=ids_to_delete)
        return True
    return False